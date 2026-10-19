"""backend/tests/test_exportar_docx.py — exportar a Word lo que el editor guarda (2026-10-19). Pruebas
puras: se genera el .docx en memoria y se lee con python-docx. Textos INVENTADOS.
Uso: cd backend && .venv/bin/python -m pytest -q tests/test_exportar_docx.py
"""
import io
import os

os.environ.setdefault("DATABASE_URL", "postgresql+asyncpg://x:y@127.0.0.1/z")
os.environ.setdefault("JWT_SECRET", "prueba")

from docx import Document  # noqa: E402

from app.servicios import exportar_docx as x  # noqa: E402

CUERPO = {"type": "doc", "content": [
    {"type": "paragraph", "attrs": {"textAlign": "justify"}, "content": [
        {"type": "text", "text": "Es grato "}, {"type": "text", "text": "dirigirme", "marks": [{"type": "bold"}]},
        {"type": "text", "text": " a usted"}, {"type": "footnote", "attrs": {"texto": "Conforme al artículo 159."}}, {"type": "text", "text": "."}]},
    {"type": "heading", "attrs": {"level": 2}, "content": [{"type": "text", "text": "Antecedentes"}]},
    {"type": "bulletList", "content": [{"type": "listItem", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Primero"}]}]},
                                       {"type": "listItem", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Segundo"}]}]}]},
    {"type": "table", "content": [{"type": "tableRow", "content": [{"type": "tableHeader", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "A"}]}]},
                                                                    {"type": "tableCell", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "B"}]}]}]}]},
    {"type": "saltoPagina"},
    {"type": "paragraph", "content": [{"type": "text", "text": "Fin."}]},
]}
CAMPOS = {"lugar": "Lima", "fecha": "2026-10-18", "destinatario_tratamiento": "Señor:", "destinatario_nombre": "JUAN PRUEBA",
          "destinatario_cargo": "Alcalde", "asunto": "Pedido de información", "referencia": "Oficio N.° 0045-2026", "con_copia": "Presidencia\nComisión",
          "iniciales": "EAC/jp", "fmt_margenes": "2", "fmt_tamano": "12"}


def _textos(data: bytes) -> list[str]:
    d = Document(io.BytesIO(data))
    return [p.text for p in d.paragraphs]


def test_documento_a_docx_lleva_encabezado_cuerpo_notas_y_cierre():
    data = x.documento_a_docx(CAMPOS, CUERPO, titulo="Prueba", tipo_nombre="Oficio", numero_completo="")
    t = _textos(data)
    assert t[0] == "OFICIO N.° ___" and t[1] == "Lima, 18 de octubre de 2026"
    assert "Señor:" in t and "JUAN PRUEBA" in t and "Presente.-" in t
    assert any(p.startswith("ASUNTO: Pedido") for p in t) and any(p.startswith("REFERENCIA: Oficio") for p in t)
    assert any(p.startswith("Es grato dirigirme a usted1.") for p in t)       # la llamada de nota, en superíndice
    assert "Antecedentes" in t and "Primero" in t and "Segundo" in t and "Fin." in t
    assert "Atentamente," in t and "c.c.: Presidencia" in t and "EAC/jp" in t
    assert "Notas" in t and any("1 Conforme al artículo 159." == p for p in t)
    d = Document(io.BytesIO(data))
    assert d.tables[0].cell(0, 0).text == "A" and d.tables[0].cell(0, 1).text == "B"
    assert abs(d.sections[0].left_margin.cm - 2) < 0.01 and d.styles["Normal"].font.size.pt == 12


def test_partes_a_docx_y_nombre():
    data = x.partes_a_docx("Proyecto de prueba", [("Fórmula legal", CUERPO), ("Exposición de motivos", None)])
    t = _textos(data)
    assert t[0] == "Proyecto de prueba" and "Fórmula legal" in t and "Exposición de motivos" in t
    assert x.nombre_docx("OF-0012-2026", "Pedido de información sobre agua") == "OF-0012-2026_Pedido-de-información-sobre-agua.docx"
    assert x.nombre_docx("", "", "Oficio") == "Oficio.docx"


def test_control_de_cambios_sale_como_w_ins_y_w_del():
    cuerpo = {"type": "doc", "content": [{"type": "paragraph", "content": [
        {"type": "text", "text": "Hola "},
        {"type": "text", "text": "mundo", "marks": [{"type": "insercion", "attrs": {"autor": "Ana", "en": "2026-10-19T10:00:00.000Z"}}]},
        {"type": "text", "text": " adiós", "marks": [{"type": "supresion", "attrs": {"autor": "Luis", "en": "2026-10-19T10:01:00.000Z"}}]},
    ]}]}
    data = x.documento_a_docx({}, cuerpo)
    d = Document(io.BytesIO(data))
    xml = d.element.xml
    assert '<w:ins ' in xml and 'w:author="Ana"' in xml and 'w:date="2026-10-19T10:00:00Z"' in xml
    assert '<w:del ' in xml and 'w:author="Luis"' in xml and '<w:delText' in xml
    # python-docx solo lee los runs directos: «mundo» va dentro de w:ins y « adiós» como w:delText.
    assert d.paragraphs[0].text == "Hola " and "<w:t>mundo</w:t>" in xml and "adiós</w:delText>" in xml


def test_color_de_texto_viaja_al_word():
    # Color de texto (2026-10-19): la marca del editor sale como color del run; un hex raro, no.
    cuerpo = {"type": "doc", "content": [{"type": "paragraph", "content": [
        {"type": "text", "text": "rojo", "marks": [{"type": "colorTexto", "attrs": {"color": "a6122e"}}]},
        {"type": "text", "text": " nada", "marks": [{"type": "colorTexto", "attrs": {"color": "zzz"}}]},
    ]}]}
    d = Document(io.BytesIO(x.documento_a_docx({}, cuerpo)))
    runs = d.paragraphs[0].runs
    assert str(runs[0].font.color.rgb) == "A6122E" and runs[1].font.color.rgb is None


def test_markdown_del_cuerpo_y_del_documento():
    from app.servicios import exportar_md as md

    texto = md.documento_a_md(CAMPOS, CUERPO, titulo="Prueba", numero_completo="OF-0012-2026")
    assert texto.startswith("# OF-0012-2026\n")
    assert "**Asunto:** Pedido de información" in texto
    assert "Es grato **dirigirme** a usted[^1]." in texto
    assert "## Antecedentes" in texto and "- Primero" in texto and "- Segundo" in texto
    assert "| A | B |" in texto and "|---|---|" in texto
    assert "\n---\n" in texto and "[^1]: Conforme al artículo 159." in texto
    assert md.partes_a_md("PL", [("Fórmula", {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "x"}]}]})]).startswith("# PL\n\n## Fórmula\n\nx\n")


def test_formato_sugerido_no_viaja_al_word():
    # Control de cambios, segunda parte (2026-10-19): una negrita o un centrado aún no aceptados
    # (marca `cambioFormato`, atributos `cambio` del párrafo) salen al Word como estaban: sin formato
    # y sin cambio registrado; solo w:ins / w:del llevan sugerencias.
    cuerpo = {"type": "doc", "content": [{"type": "paragraph", "attrs": {"textAlign": "justify", "cambio": "insercion", "cambioAutor": "Ana", "cambioEn": "2026-10-19T10:00:00Z"}, "content": [
        {"type": "text", "text": "pendiente", "marks": [{"type": "cambioFormato", "attrs": {"autor": "Ana", "en": "", "antes": '{"bold":false}', "despues": '{"bold":true,"textAlign":"center"}'}}]},
    ]}]}
    d = Document(io.BytesIO(x.documento_a_docx({}, cuerpo)))
    xml = d.element.xml
    run = d.paragraphs[0].runs[0]
    assert run.text == "pendiente" and run.bold is None and d.paragraphs[0].alignment in (None, x.ALINEACION["justify"])
    assert "<w:ins " not in xml and "<w:del " not in xml
